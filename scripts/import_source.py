"""把玩家整理的数据源 Excel 转成 data/characters.csv、data/classes.csv。

数据源格式：两个 sheet「01-全可加入角色信息」「02-兵种职业信息」（见 docs/固定成长养成规划器.md）。
用法：python scripts/import_source.py 数据源.xlsx
"""
import csv
import os
import re
import sys

from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
STATS = ["HP", "力量", "魔力", "速度", "技巧", "防守", "魔防", "幸运", "魅力"]
SKILLS = ["剑术", "枪术", "斧术", "弓术", "格斗术", "黑魔法", "白魔法", "指挥", "步兵", "马术", "重装", "飞行"]
# 职业「可使用的技能」11 列的顺序（没有指挥）
USE_KEYS = ["剑术", "枪术", "斧术", "弓术", "格斗术", "黑魔法", "白魔法", "步兵术", "马术", "重装术", "飞行术"]
TIER_MIN = {"基础": 1, "初级": 5, "中级": 20, "上级": 35, "最上级": 45, "神将": 45}
TIER_NAME = {"基础": "基本职", "初级": "初级职", "中级": "中级职", "上级": "上级职", "最上级": "最上级职", "神将": "神将职"}
# 已知的精通被动属性加成（数据源没有结构化字段，按玩家反馈补充）
MASTERY = {"天翼兵": ("魔防", 3)}
# 自动推荐默认「靠后」/「不参与」的职业
LOW = {"舞者", "神乐神将"}
OFF = {"战象兵"}
NORM = [("箭术", "弓术"), ("黑魔术", "黑魔法"), ("白魔术", "白魔法")]
REQ_KEYS = ["剑术", "枪术", "斧术", "弓术", "格斗术", "黑魔法", "白魔法", "马术", "重装术", "飞行术"]
RANK_RE = r"(S\+|A\+|B\+|C\+|D\+|E\+|[SABCDEF])"


def clean(v):
    if v is None:
        return ""
    s = str(v).strip()
    return "" if s in ("None",) else s


def norm_skill(s):
    for a, b in NORM:
        s = s.replace(a, b)
    s = re.sub(r"重装(?!术)", "重装术", s)
    s = re.sub(r"飞行(?!术)", "飞行术", s)
    return s


def parse_req(text):
    """「剑术/斧术任意一项C&重装术E+、道具…」→ (主要技能, 选择技能)，格式如「马术D、弓术B」。"""
    head = re.split(r"、道具|，|,|（|\n|/解锁", text)[0]
    head = norm_skill(head)
    # 「剑术A+箭术C」里的 + 是连接符（后面紧跟技能名）
    head = re.sub(r"\+(?=[剑枪斧弓格黑白马重飞])", "&", head)
    main, sel = [], []
    for part in head.split("&"):
        part = part.strip()
        m = re.match(r"(.+?)(?:任意一项|任一)" + RANK_RE, part)
        if m:
            for sk in m.group(1).split("/"):
                sk = norm_skill(sk.strip())
                if sk in REQ_KEYS:
                    sel.append(sk + m.group(2))
            continue
        m = re.match(r"(.+?)(?:至少达到)?" + RANK_RE + r"$", part)
        if m and norm_skill(m.group(1).strip()) in REQ_KEYS:
            main.append(norm_skill(m.group(1).strip()) + m.group(2))
    return "、".join(main) or "—", "、".join(sel) or "—"


def import_chars(ws):
    hdr = (["编号", "所属势力", "角色", "性别", "初始兵种"] + STATS + ["成长合计", "初始等级"]
           + ["初始" + s for s in STATS] + SKILLS
           + ["擅长", "弱项", "个人特技", "个人特技说明", "血印", "血印说明", "最早可加入章节", "加入条件", "挖角条件",
              "装备特技", "可习得战技"])
    rows, faction = [], ""
    for r in range(4, ws.max_row + 1):
        g = lambda c: ws.cell(r, c).value
        name = clean(g(5)).split("\n")[0]
        if not name:
            continue
        if clean(g(3)):
            faction = clean(g(3)).replace("\n", "")
        elif not clean(g(2)):
            faction = ""  # 无编号的角色（数据不全）不继承上一行的势力
        growth = [g(c) for c in range(18, 27)]
        if all(v in (None, "") for v in growth):
            continue  # 无成长率数据的 NPC 不收录
        ranks, good, weak = [], [], []
        for i, c in enumerate(range(39, 51)):
            v = clean(g(c))
            tag = "擅长" if "擅长" in v else "弱项" if "弱项" in v else ""
            rk = v.replace("擅长", "").replace("弱项", "").replace("\n", "").strip()
            if rk in ("", "--"):
                rk = "E" if v else "—"   # 「--」= 初始E（数据源不特意标记）
            if tag == "擅长":
                good.append(SKILLS[i])
            elif tag == "弱项":
                weak.append(SKILLS[i])
            ranks.append(rk)
        lv = g(28)
        init = [g(c) for c in range(29, 38)]
        join_cells = lambda cs: "；".join(clean(g(c)).replace("\n", "") for c in cs if clean(g(c)))
        rows.append([clean(g(2)), faction, name, clean(g(7)) or "—", clean(g(6)) or "—"]
                    + [clean(v) for v in growth] + [""]
                    + [clean(lv)] + [clean(v) for v in init] + ranks
                    + ["、".join(good), "、".join(weak), clean(g(9)), clean(g(10)), clean(g(13)), clean(g(14)),
                       clean(g(15)), clean(g(16)), clean(g(17)), join_cells(range(51, 57)), join_cells(range(57, 71))])
    return hdr, rows


def import_classes(ws):
    hdr = (["序号", "阶级", "兵种名", "最低转职等级", "限定"] + STATS + ["精通加成属性", "精通加成值"]
           + ["基础" + s for s in STATS]
           + ["转职·主要技能", "转职·选择技能", "转职条件", "使用武器", "技能EXP加成", "移动力", "兵种特性",
              "兵种固有特技", "精通所需EXP", "精通特技", "纳入推荐"])
    rows = []
    for r in range(4, ws.max_row + 1):
        g = lambda c: ws.cell(r, c).value
        name = clean(g(3))
        if not name:
            continue
        tier_raw = clean(g(4))
        tier = re.split(r"[/\n]", tier_raw)[0].strip()
        restrict = "女" if "女性专用" in tier_raw else ""
        m = re.search(r"（(.+?)限定）", tier_raw)
        if m:
            restrict = m.group(1)
        cond = clean(g(5))
        m = re.search(r"推荐LV(\d+)", cond)
        min_lv = int(m.group(1)) if m else TIER_MIN.get(tier, 1)
        main, sel = parse_req(cond) if cond not in ("", "--") else ("—", "—")
        use = "・".join(k for k, c in zip(USE_KEYS, range(34, 45)) if clean(g(c)) == "〇")
        mb = MASTERY.get(name, ("", ""))
        avail = "靠后" if name in LOW else "否" if name in OFF else "是"
        rows.append([str(len(rows) + 1), TIER_NAME.get(tier, tier), name, str(min_lv), restrict]
                    + [clean(g(c)) for c in range(8, 17)] + [mb[0], str(mb[1])]
                    + [clean(g(c)) for c in range(18, 27)]
                    + [main, sel, cond.replace("\n", "；"), use, clean(g(6)), clean(g(7)), clean(g(28)).replace("\n", "/"),
                       "".join(clean(g(c)) for c in range(29, 32) if clean(g(c)) not in ("", "--")),
                       clean(g(32)), clean(g(33)).replace("\n", ""), avail])
    return hdr, rows


def write(name, hdr, rows):
    with open(os.path.join(ROOT, "data", name), "w", newline="", encoding="utf-8-sig") as fp:
        w = csv.writer(fp)
        w.writerow(hdr)
        w.writerows(rows)
    print(f"data/{name}: {len(rows)} 行")


def main(path):
    wb = load_workbook(path, data_only=True)
    write("characters.csv", *import_chars(wb["01-全可加入角色信息"]))
    write("classes.csv", *import_classes(wb["02-兵种职业信息"]))


if __name__ == "__main__":
    main(sys.argv[1])
