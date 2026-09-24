"""生成《固定成长养成规划器》Excel。

数据源：data/characters.csv（角色成长率）、data/classes.csv（职业补正）
输出：tools/固定成长养成规划器.xlsx

用法：python scripts/build_growth_planner.py
"""
import csv
import os

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "tools", "固定成长养成规划器.xlsx")

STATS = ["HP", "力量", "魔力", "速度", "技巧", "防守", "魔防", "幸运", "魅力"]
NS = len(STATS)
DATA_ROWS = 150          # 数据表预留行数（第3行起），方便以后追加角色/职业
DATA_LAST = 2 + DATA_ROWS
COMBO_ROWS = 6000        # 推荐路线搜索的组合上限
REF_ROWS = 4000          # 中途换职优化的候选上限
TL_ROWS = 100            # 逐级成长表行数
# 武器熟练度：角色表列名 / 在职业表文本里搜索用的关键字
SKILLS = ["剑术", "枪术", "斧术", "弓术", "格斗术", "黑魔法", "白魔法", "指挥", "步兵", "马术", "重装", "飞行"]
SKILL_KEYS = ["剑术", "枪术", "斧术", "弓术", "格斗术", "黑魔法", "白魔法", "指挥", "步兵术", "马术", "重装术", "飞行术"]
NK = len(SKILLS)
RANKS = ["F", "E", "E+", "D", "D+", "C", "C+", "B", "B+", "A", "A+", "S", "S+"]

SH_MAIN, SH_CHAR, SH_CLASS = "养成规划", "角色成长率", "职业"
CHAR_Q, CLASS_Q = f"'{SH_CHAR}'", f"'{SH_CLASS}'"

# ---------- 样式 ----------
FONT = "Microsoft YaHei"
f_base = Font(name=FONT, size=10)
f_bold = Font(name=FONT, size=10, bold=True)
f_title = Font(name=FONT, size=16, bold=True, color="1F3864")
f_sec = Font(name=FONT, size=12, bold=True, color="FFFFFF")
f_hdr = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_input = Font(name=FONT, size=10, color="0000FF", bold=True)
f_note = Font(name=FONT, size=9, color="595959")
f_warn = Font(name=FONT, size=10, bold=True, color="C00000")
fill_sec = PatternFill("solid", fgColor="1F3864")
fill_hdr = PatternFill("solid", fgColor="2F5597")
fill_input = PatternFill("solid", fgColor="FFF2CC")
fill_info = PatternFill("solid", fgColor="DEEBF7")
fill_grey = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="center", wrap_text=True)
left_nowrap = Alignment(horizontal="left", vertical="center")


def num(v):
    try:
        f = float(v)
        return int(f) if f.is_integer() else f
    except (TypeError, ValueError):
        return v


def load_csv(name):
    with open(os.path.join(ROOT, "data", name), encoding="utf-8-sig") as fp:
        rows = list(csv.reader(fp))
    return rows[0], rows[1:]


def style_range(ws, ref, font=None, fill=None, align=None, brd=True):
    for row in ws[ref]:
        for c in row:
            if font:
                c.font = font
            if fill:
                c.fill = fill
            if align:
                c.alignment = align
            if brd:
                c.border = border


def section(ws, row, text, last_col="M"):
    ws.merge_cells(f"A{row}:{last_col}{row}")
    c = ws[f"A{row}"]
    c.value = text
    c.font = f_sec
    c.fill = fill_sec
    c.alignment = left_nowrap
    ws.row_dimensions[row].height = 22


def header(ws, row, values, start_col=1):
    for i, v in enumerate(values):
        c = ws.cell(row, start_col + i, v)
        c.font = f_hdr
        c.fill = fill_hdr
        c.alignment = center
        c.border = border


# ======================================================================
# 数据表
# ======================================================================
def build_data_sheet(ws, title, hdr, rows, front, growth_cols, extra_formula=None, widths=None):
    """front：放在成长率之前的列名；growth_cols：9 个成长率列名；其余列原样追加在后面。"""
    rest = [h for h in hdr if h not in front and h not in growth_cols]
    cols = front + growth_cols + (["成长合计"] if extra_formula else []) + [h for h in rest if h != "成长合计"]
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=12)
    ws["A1"] = title
    ws["A1"].font = Font(name=FONT, size=14, bold=True, color="1F3864")
    header(ws, 2, cols)
    ws.row_dimensions[2].height = 32
    idx = {h: i for i, h in enumerate(hdr)}
    for r, row in enumerate(rows, start=3):
        for c, h in enumerate(cols, start=1):
            if h == "成长合计" and extra_formula:
                v = extra_formula(r)
            else:
                v = row[idx[h]] if h in idx else ""
                v = num(v) if h in growth_cols or h in ("序号", "最低转职等级", "移动力") else v
                if v == "":
                    v = None
            cell = ws.cell(r, c, v)
            cell.font = f_base
            cell.border = border
            cell.alignment = center if (h in growth_cols or h in front or h in ("成长合计", "移动力")) else left
    g0 = len(front) + 1
    for r in range(3, 3 + len(rows)):
        for c in range(g0, g0 + NS):
            ws.cell(r, c).fill = fill_info
    ws.freeze_panes = ws.cell(3, len(front) + 1)
    ws.auto_filter.ref = f"A2:{CL(len(cols))}{2 + len(rows)}"
    for i, h in enumerate(cols, start=1):
        w = (widths or {}).get(h)
        if w is None:
            w = 7 if h in growth_cols else 14
        ws.column_dimensions[CL(i)].width = w
    return cols


def build_char_sheet(wb):
    hdr, rows = load_csv("characters.csv")
    ws = wb.create_sheet(SH_CHAR)
    front = ["序号", "阵营", "角色", "日文名", "初始兵种"]
    widths = {"序号": 6, "阵营": 14, "角色": 12, "日文名": 14, "初始兵种": 10, "成长合计": 9,
              "职业偏向（图鉴·战斗特征）": 36, "个人技能": 36, "说明": 60, "血印1": 30, "血印2": 30}
    cols = build_data_sheet(ws, "角色成长率（个人成长率 %，可直接在此订正；角色名需与下拉框一致）", hdr, rows, front, STATS,
                            extra_formula=lambda r: f'=IF(COUNT(F{r}:N{r})=0,"—",SUM(F{r}:N{r}))', widths=widths)
    assert cols[2] == "角色" and cols[5:14] == STATS
    return len(rows), cols


def build_class_sheet(wb):
    hdr, rows = load_csv("classes.csv")
    ws = wb.create_sheet(SH_CLASS)
    front = ["序号", "阶级", "兵种名", "最低转职等级", "纳入推荐"]
    widths = {"序号": 6, "阶级": 9, "兵种名": 13, "最低转职等级": 9, "纳入推荐": 8, "转职条件": 22,
              "使用武器": 26, "中文说明": 50, "解放条件": 30, "代表角色": 24}
    cols = build_data_sheet(ws, "职业（职业成长率补正 %，可直接在此订正/追加）", hdr, rows, front, STATS, widths=widths)
    assert cols[2] == "兵种名" and cols[5:14] == STATS
    ws.cell(2, 5).value = "纳入推荐\n(是/否/靠后)"
    for r in range(3, 3 + len(rows)):
        for c in (4, 5):
            ws.cell(r, c).fill = fill_input
            ws.cell(r, c).font = f_input
    dv = DataValidation(type="list", formula1='"是,靠后,否"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"E3:E{DATA_LAST}")
    note_r = 3 + len(rows) + 1
    ws.cell(note_r, 1, "说明：").font = f_bold
    notes = [
        "「最低转职等级」：初级5、中级20、上级35、最上级45（暂定）、神将45（暂定）。自动推荐只会在达到该等级后才转入此职业。",
        "「纳入推荐」：是=参与推荐；靠后=参与推荐，但同等条件下排在其他路线后面（舞者、飞天女神将）；否=不参与自动推荐（手动路线仍可选）。没解锁或不想用的职业可改成否。",
        "「转职·主要技能」为转职必需的武器熟练度（全部满足），「转职·选择技能」为满足其一即可；规划页会据此检查路线可行性。「使用武器」决定在该职业期间能练哪些熟练度。",
        "要追加职业，在表格末尾接着填一行即可（第152行以内），兵种名会自动出现在养成规划的下拉框里。",
    ]
    for i, t in enumerate(notes, start=1):
        ws.merge_cells(start_row=note_r + i, start_column=1, end_row=note_r + i, end_column=20)
        c = ws.cell(note_r + i, 1, t)
        c.font = f_note
        c.alignment = left_nowrap
    return len(rows), cols


# ======================================================================
# 主页面
# ======================================================================
SC = [CL(4 + j) for j in range(NS)]           # 主页面属性列 D..L
WC = [CL(2 + s) for s in range(NK)]           # 主页面熟练度列 B..M
DC = [CL(6 + j) for j in range(NS)]           # 数据表属性列 F..N
LAST_VIS = "P"                                # 可见区最右列

CHAR_NAMES = f"{CHAR_Q}!$C$3:$C${DATA_LAST}"
CLASS_NAMES = f"{CLASS_Q}!$C$3:$C${DATA_LAST}"
H0 = 5                                        # 辅助表首行
HZ = DATA_ROWS + 1                            # 职业辅助表的「空行」序号：序号0统一映射到这一行（全为0）
CLS_LAST = H0 + DATA_ROWS                     # 职业辅助表末行（含空行）

# 可见区行号
R_WSEC, R_WHDR, R_WINIT, R_WCUR, R_WUSE = 17, 18, 19, 20, 21
R_BSEC, R_BHDR, R_B0 = 23, 24, 25
R_OSEC, R_OHDR, R_O0 = 31, 32, 33
R_SSEC, R_SHDR, R_S0 = 39, 40, 41
R_FSEC, R_FHDR, R_F0 = 48, 49, 50
R_MSEC, R_MHDR, R_M0 = 55, 56, 57
R_TSEC, R_THDR, R_T0 = 64, 65, 66
R_TLAST = R_T0 + TL_ROWS - 1


def class_col(letter, last=DATA_LAST + 1):
    return f"{CLASS_Q}!${letter}$3:${letter}${last}"


def Z(x):
    """职业序号 0（找不到）→ 指向全 0 的空行，避免 INDEX(...,0) 返回整列。"""
    return f"IF({x}=0,{HZ},{x})"


class Alloc:
    def __init__(self, start):
        self.c = start

    def take(self, n=1):
        cols = [CL(self.c + i) for i in range(n)]
        self.c += n
        return cols if n > 1 else cols[0]


def rng(col, r1, r2):
    return f"${col}${r1}:${col}${r2}"


def mrng(c1, c2, r1, r2):
    return f"${c1}${r1}:${c2}${r2}"


def build_main(wb, char_cols, class_cols):
    ws = wb.active
    ws.title = SH_MAIN
    ws.sheet_view.showGridLines = False
    for row in ws.iter_rows(min_row=1, max_row=R_TLAST, max_col=16):
        for c in row:
            c.font = f_base
    widths = {"A": 12, "B": 10, "C": 13, "M": 50, "N": 24, "O": 40, "P": 40}
    for col in "ABCDEFGHIJKLMNOP":
        ws.column_dimensions[col].width = widths.get(col, 8)

    # 数据表中的列字母
    ccol = {h: CL(i + 1) for i, h in enumerate(char_cols)}
    kcol = {h: CL(i + 1) for i, h in enumerate(class_cols)}
    SK_CHAR = [ccol[s] for s in SKILLS]
    C_USE, C_MAIN, C_SEL = kcol["使用武器"], kcol["转职·主要技能"], kcol["转职·选择技能"]
    C_COND, C_UNLOCK = kcol["转职条件"], kcol["解放条件"]

    # ------------------------------------------------------------------
    # 辅助区列分配（Q 列起全部隐藏）
    # ------------------------------------------------------------------
    A = Alloc(18)
    SCL_L, SCL_V = A.take(), A.take()
    W_LO, W_HI, W_E, W_N = A.take(), A.take(), A.take(), A.take()
    V9 = A.take(NS)            # 行5 C0，行6 C4，行7 K，行8 KT
    V13 = A.take(13)           # 行5 等级表，行6 当前熟练度，行7 可练标记，行8~12 推荐1各阶段累计可练
    CLS = {"name": A.take(), "min": A.take(), "avail": A.take(), "low": A.take(),
           "G": A.take(NS), "T": A.take(NK), "RM": A.take(NK), "RS": A.take(NK), "NM": A.take(NK), "SU": A.take(NK),
           "SA": A.take(), "SS": A.take(), "EL": A.take(3), "CN": A.take(3), "ELU": A.take(4), "CNU": A.take(4)}
    FL = [{"idx": A.take(), "name": A.take(), "C": A.take(NS)} for _ in range(3)]
    UL = A.take(4)
    CMB = {k: A.take() for k in ("i1", "i2", "i3", "id1", "id2", "id3", "bad", "low", "met", "short", "total", "score")}
    REF = {k: A.take() for k in ("k", "j", "x", "B", "valid", "bad", "low", "met", "short", "total", "score", "best")}
    RK = {k: A.take() for k in ("pos", "valid", "met", "short", "total", "bad", "low", "i1", "i2", "i3",
                                "k", "x", "B", "P", "head", "split", "first", "second")}
    RK["FROM"], RK["IDX"], RK["NAME"], RK["N"] = A.take(6), A.take(6), A.take(6), A.take(6)
    SEG = {k: A.take() for k in ("from", "idx", "name", "n", "prev", "piece", "scan")}
    SEG["U"] = A.take(NK)
    MAN = {k: A.take() for k in ("from", "idx", "name", "n", "prev", "piece", "scan")}
    MAN["U"] = A.take(NK)
    TLH = {"idx": A.take(), "cum": A.take(NS)}
    LAST_HELPER = A.c - 1

    H = ws.__setitem__
    # 窗口表：k=0..4 → 行 5..9
    E = [f"${W_E}${H0 + k}" for k in range(5)]
    N = [f"${W_N}${H0 + k}" for k in range(5)]
    LO = [f"${W_LO}${H0 + k}" for k in range(5)]
    EW = rng(W_E, H0 + 1, H0 + 4)        # k=1..4
    NW = rng(W_N, H0 + 1, H0 + 4)

    # 标量
    scal = {}
    srow = [H0]

    def S(name, formula):
        r = srow[0]
        H(f"{SCL_L}{r}", name)
        H(f"{SCL_V}{r}", formula)
        scal[name] = f"${SCL_V}${r}"
        srow[0] += 1
        return scal[name]

    def ref(name):
        return scal[name]

    # 预先确定标量位置（公式里互相引用）
    names_order = (["startIdx", "tgtIdx", "tgtMin", "force", "cnt1", "cnt2", "cnt3", "m1", "m2", "m3", "M", "listOK",
                    "tgtCnt", "mm1", "mm2", "mm3", "mm4", "sz1", "sz2", "sz3", "sz4", "O1", "O2", "O3", "O4", "R",
                    "NN1", "NN2", "NN3", "NN4", "head1", "head2", "head3", "head4", "P1", "P2", "P3", "P4",
                    "validT", "badT", "lowT", "entryLv", "sel", "rankMissing", "manualBad"])
    for i, nm in enumerate(names_order):
        scal[nm] = f"${SCL_V}${H0 + i}"

    NAMES = rng(CLS["name"], H0, CLS_LAST)
    MINC = rng(CLS["min"], H0, CLS_LAST)
    LOWC = rng(CLS["low"], H0, CLS_LAST)
    SAC, SSC = rng(CLS["SA"], H0, CLS_LAST), rng(CLS["SS"], H0, CLS_LAST)
    GMAT = mrng(CLS["G"][0], CLS["G"][-1], H0, CLS_LAST)
    TMAT = mrng(CLS["T"][0], CLS["T"][-1], H0, CLS_LAST)
    NMMAT = mrng(CLS["NM"][0], CLS["NM"][-1], H0, CLS_LAST)
    SUMAT = mrng(CLS["SU"][0], CLS["SU"][-1], H0, CLS_LAST)
    RANKROW = mrng(V13[0], V13[12], H0, H0)
    CURROW = mrng(V13[0], V13[NK - 1], H0 + 1, H0 + 1)
    CTROW = mrng(V13[0], V13[NK - 1], H0 + 2, H0 + 2)
    UCUM = mrng(V13[0], V13[NK - 1], H0 + 3, H0 + 7)
    KROW = mrng(V9[0], V9[-1], H0 + 2, H0 + 2)
    KTROW = mrng(V9[0], V9[-1], H0 + 3, H0 + 3)
    START = "$D$9:$L$9"
    TGT = "$D$10:$L$10"
    FLC = [mrng(FL[k]["C"][0], FL[k]["C"][-1], H0, H0 + DATA_ROWS - 1) for k in range(3)]
    FLI = [rng(FL[k]["idx"], H0, H0 + DATA_ROWS - 1) for k in range(3)]
    ULMAT = mrng(UL[0], UL[3], H0, H0 + DATA_ROWS - 1)
    OROW = f"{ref('O1')}:{ref('O4')}"
    NNROW = f"{ref('NN1')}:{ref('NN4')}"
    HEADROW = f"{ref('head1')}:{ref('head4')}"
    PROW = f"{ref('P1')}:{ref('P4')}"
    trow = lambda x: f"INDEX({TMAT},{Z(x)},0)"

    def badc(x, u):
        """转入职业 x 时，熟练度无法满足的项数（u = 之前职业累计可练向量表达式）。"""
        zx = Z(x)
        return (f"(SUMPRODUCT(INDEX({NMMAT},{zx},0)*((({u})*{CTROW})=0))"
                f"+IF(AND(INDEX({SAC},{zx})=1,INDEX({SSC},{zx})=0,"
                f"SUMPRODUCT(INDEX({SUMAT},{zx},0)*((({u})*{CTROW})>0))=0),1,0))")

    def score(ok, bad, low, met, short, total, i):
        # 量级控制在 14 位有效数字内（LibreOffice 只保留 15 位），否则区分并列用的小数会丢失
        return (f"=IF({ok},(9-MIN({bad},9))*1E+9+{met}*1E+8+({low}=0)*1E+7-MIN({short},999)*1E+4"
                f"+MIN({total},9999)-{i}/1E+4,-1E+12-{i})")

    def met_short_total(fin):
        return (f'SUMPRODUCT(({TGT}<>"")*({fin}>={TGT}))',
                f'SUMPRODUCT(({TGT}<>"")*(({TGT}-{fin})>0)*({TGT}-{fin}))',
                f"SUMPRODUCT({fin})")

    # ------------------------------------------------------------------
    # 标题 / 输入区
    # ------------------------------------------------------------------
    ws.merge_cells(f"A1:{LAST_VIS}1")
    ws["A1"] = "火焰纹章 万缕千丝 ｜ 固定成长 · 养成路线规划器"
    ws["A1"].font = f_title
    ws.row_dimensions[1].height = 30
    ws.merge_cells(f"A2:{LAST_VIS}2")
    ws["A2"] = ("用法：黄底蓝字是输入格。① 选角色 → 填起点/目标的等级、职业、属性（目标属性不填=不作要求）→ ② 填当前武器熟练度 → "
                "③ 看推荐路线与中途换职优化，在「显示路线」切换 → ④ 查看每段的转职等级、所需熟练度与解锁条件 → ⑥ 逐级成长表。"
                "规则：每升1级，各属性经验槽 +（角色成长率+职业补正，最低0），满100则该属性+1。")
    ws["A2"].font = f_note
    ws["A2"].alignment = left
    ws.row_dimensions[2].height = 42

    ws["A4"] = "① 角色"
    ws["A4"].font = f_bold
    ws.merge_cells("B4:C4")
    ws["B4"] = "蕾达"
    style_range(ws, "B4:C4", font=f_input, fill=fill_input, align=center)
    char_match = f"MATCH($B$4,{CHAR_NAMES},0)"
    ws["E4"] = "阵营"
    ws.merge_cells("F4:G4")
    ws["F4"] = f'=IFERROR(INDEX({CHAR_Q}!$B$3:$B${DATA_LAST},{char_match}),"")'
    ws["H4"] = "初始兵种"
    ws.merge_cells("I4:J4")
    ws["I4"] = f'=IFERROR(INDEX({CHAR_Q}!$E$3:$E${DATA_LAST},{char_match}),"")'
    ws.merge_cells("K4:M4")
    ws["K4"] = (f'=IF($B$4="","请选择角色",IF(ISNA({char_match}),"⚠ 角色表里找不到该角色",'
                f'IF(COUNT(INDEX({CHAR_Q}!$F$3:$N${DATA_LAST},{char_match},0))=0,"⚠ 该角色成长率数据缺失（按0计算）","✓ 成长率已读取")))')
    for r in ("E4", "H4"):
        ws[r].font = f_bold
        ws[r].alignment = center
    style_range(ws, "F4:G4", fill=fill_grey, align=center)
    style_range(ws, "I4:J4", fill=fill_grey, align=center)
    ws["K4"].font = f_bold
    ws["K4"].alignment = left_nowrap

    header(ws, 6, ["项目", "等级", "职业"] + STATS + ["说明"])
    labels = {7: "角色成长率%", 8: "起点职业补正%", 9: "起点", 10: "目标", 11: "需提升", 12: "每级需成长%"}
    notes = {7: "来自「角色成长率」表", 8: "来自「职业」表，仅供参考",
             9: "填当前等级、职业和面板属性（示例数据，请替换）", 10: "目标属性留空=不作要求；目标职业留空=45级后沿用上一职业",
             11: "目标 − 起点", 12: "达成目标平均每级需要的（角色+职业）成长，可对照职业补正"}
    for r, t in labels.items():
        ws.cell(r, 1, t).font = f_bold
        ws.cell(r, 1).alignment = center
        ws.cell(r, 1).border = border
        ws.cell(r, 13, notes[r]).font = f_note
        ws.cell(r, 13).alignment = left_nowrap
        ws.cell(r, 13).border = border
    for j in range(NS):
        s, d = SC[j], DC[j]
        ws[f"{s}7"] = f"=IFERROR(INDEX({CHAR_Q}!{d}$3:{d}${DATA_LAST},{char_match})+0,0)"
        ws[f"{s}8"] = f"=IFERROR(INDEX({class_col(d)},MATCH($C$9,{CLASS_NAMES},0))+0,0)"
        ws[f"{s}11"] = f'=IF(OR({s}10="",{s}9=""),"",{s}10-{s}9)'
        ws[f"{s}12"] = (f'=IF(OR({s}10="",$B$10="",$B$9="",$B$10<=$B$9),"",'
                        f'ROUND(MAX(0,({s}10-{s}9)*100-$B$14)/($B$10-$B$9),1))')
    for r in (7, 8, 11, 12):
        style_range(ws, f"B{r}:L{r}", fill=fill_grey, align=center)
    ws["B9"], ws["C9"] = 10, "猎兵"
    ws["B10"], ws["C10"] = 40, "狙击手"
    start = [26, 10, 6, 14, 11, 7, 6, 5, 9]
    target = [40, 22, None, 30, 27, None, None, None, None]
    for j in range(NS):
        ws[f"{SC[j]}9"] = start[j]
        ws[f"{SC[j]}10"] = target[j]
    style_range(ws, "B9:L10", font=f_input, fill=fill_input, align=center)

    # 参数
    ws["A13"] = "换职最少练级"
    ws["A13"].font = f_bold
    ws["B13"] = 3
    style_range(ws, "B13:B13", font=f_input, fill=fill_input, align=center)
    ws.merge_cells("C13:L13")
    ws["C13"] = "「优化」路线中途换职时，换职前后两段各至少练这么多级（避免只练1级就换的无意义方案）"
    ws["C13"].font = f_note
    ws["C13"].alignment = left_nowrap
    ws["A14"] = "初始经验槽%"
    ws["A14"].font = f_bold
    ws["B14"] = 0
    ws.merge_cells("C14:D14")
    ws["C14"] = "推荐路线转职节点Lv"
    ws["C14"].font = f_bold
    ws["C14"].alignment = center
    for col, v in zip("EFGH", (5, 20, 35, 45)):
        ws[f"{col}14"] = v
    ws.merge_cells("I14:J14")
    ws["I14"] = "显示路线"
    ws["I14"].font = f_bold
    ws["I14"].alignment = center
    ws.merge_cells("K14:L14")
    ws["K14"] = "推荐1"
    style_range(ws, "B14:B14", font=f_input, fill=fill_input, align=center)
    style_range(ws, "E14:H14", font=f_input, fill=fill_input, align=center)
    style_range(ws, "K14:L14", font=f_input, fill=fill_input, align=center)
    ws.merge_cells(f"M14:{LAST_VIS}14")
    ws["M14"] = ("经验槽默认0=最保守。节点：推荐路线在这些等级换职业（初级5/中级20/上级35/最上级45）；"
                 "「优化」路线会在推荐1的某个区间中途再换一次职业。显示路线可选 推荐1~5 / 优化1~5 / 手动。")
    ws["M14"].font = f_note
    ws["M14"].alignment = left
    ws.row_dimensions[14].height = 30

    # ② 武器熟练度
    section(ws, R_WSEC, "② 武器熟练度（用于检查转职条件；每个职业只能练它「使用武器」里的熟练度）", LAST_VIS)
    header(ws, R_WHDR, ["项目"] + SKILLS + ["说明"])
    ws.merge_cells(f"N{R_WHDR}:{LAST_VIS}{R_WHDR}")
    for r, t, note in ((R_WINIT, "角色初始", "来自「角色成长率」表（—=数据缺失）"),
                       (R_WCUR, "当前(可填)", "填游戏里当前的熟练度，留空=用角色初始；×=无法习得"),
                       (R_WUSE, "计算用", "缺失按E计算")):
        ws.cell(r, 1, t).font = f_bold
        ws.cell(r, 1).alignment = center
        ws.cell(r, 1).border = border
        ws.merge_cells(f"N{r}:{LAST_VIS}{r}")
        ws[f"N{r}"] = note
        style_range(ws, f"N{r}:{LAST_VIS}{r}", font=f_note, align=left_nowrap)
    for s in range(NK):
        w, cc = WC[s], SK_CHAR[s]
        ws[f"{w}{R_WINIT}"] = f'=IFERROR(INDEX({CHAR_Q}!{cc}$3:{cc}${DATA_LAST},{char_match})&"","")'
        v = f'IF({w}{R_WCUR}<>"",{w}{R_WCUR},{w}{R_WINIT})'
        ws[f"{w}{R_WUSE}"] = f'=IF({v}="×","×",IFERROR(INDEX({RANKROW},1,MATCH({v},{RANKROW},0)),"E"))'
    style_range(ws, f"B{R_WINIT}:M{R_WINIT}", fill=fill_grey, align=center)
    style_range(ws, f"B{R_WCUR}:M{R_WCUR}", font=f_input, fill=fill_input, align=center)
    style_range(ws, f"B{R_WUSE}:M{R_WUSE}", fill=fill_grey, align=center)

    # ------------------------------------------------------------------
    # 辅助区：窗口表 / 标量 / 向量
    # ------------------------------------------------------------------
    H(f"{SCL_L}4", "【辅助计算区，请勿修改】")
    header(ws, 4, ["lo", "hi", "e", "n"], start_col=ws[f"{W_LO}1"].column)
    los = ["0", "$E$14", "$F$14", "$G$14", "$H$14"]
    his = ["$E$14", "$F$14", "$G$14", "$H$14", "999"]
    for k in range(5):
        r = H0 + k
        H(f"{W_LO}{r}", f"={los[k]}")
        H(f"{W_HI}{r}", f"={his[k]}")
        H(f"{W_E}{r}", f"=IFERROR(MIN(MAX($B$9,{W_LO}{r}),$B$10),0)")
        H(f"{W_N}{r}", f"=IFERROR(MAX(0,MIN($B$10,{W_HI}{r})-MAX($B$9,{W_LO}{r})),0)")

    S("startIdx", f"=IFERROR(MATCH($C$9,{NAMES},0),0)")
    S("tgtIdx", f'=IF($C$10="",0,IFERROR(MATCH($C$10,{NAMES},0),0))')
    S("tgtMin", f"=INDEX({MINC},{Z(ref('tgtIdx'))})")
    S("force", f'=IF(OR($C$10="",{N[4]}>0),0,IF({N[3]}>0,3,IF({N[2]}>0,2,IF({N[1]}>0,1,0))))')
    for k in range(3):
        S(f"cnt{k + 1}", f"=COUNTIF({rng(CLS['EL'][k], H0, CLS_LAST)},TRUE)")
    for k in range(3):
        S(f"m{k + 1}", f"=IF({N[k + 1]}=0,1,MAX(1,{ref(f'cnt{k + 1}')}))")
    S("M", f"={ref('m1')}*{ref('m2')}*{ref('m3')}")
    S("listOK", "=AND(" + ",".join(f"OR({N[k + 1]}=0,{ref(f'cnt{k + 1}')}>0)" for k in range(3)) + ")")
    S("tgtCnt", "=COUNT($D$10:$L$10)")
    for k in range(4):
        S(f"mm{k + 1}", f"=COUNTIF({rng(CLS['ELU'][k], H0, CLS_LAST)},TRUE)")
    for k in range(4):
        S(f"sz{k + 1}", f"=IF({N[k + 1]}>=2,{ref(f'mm{k + 1}')}*({N[k + 1]}-1),0)")
    S("O1", "=0")
    for k in range(1, 4):
        S(f"O{k + 1}", f"={ref(f'O{k}')}+{ref(f'sz{k}')}")
    S("R", f"={ref('O4')}+{ref('sz4')}")
    for k in range(4):
        S(f"NN{k + 1}", f"=MAX(1,{N[k + 1]}-1)")
    for k in range(3):
        S(f"head{k + 1}", f"=IF({ref('force')}={k + 1},1,0)")
    S("head4", f'=IF(AND($C$10<>"",{N[4]}>0),1,0)')
    for k in range(4):
        S(f"P{k + 1}", f"=${RK['IDX'][k + 1]}${H0}")
    S("validT", f"=${RK['valid']}${H0}")
    S("badT", f"=${RK['bad']}${H0}")
    S("lowT", f"=${RK['low']}${H0}")
    S("entryLv", f"=IF({N[4]}>0,{E[4]},IF({ref('force')}>0,INDEX({rng(W_E, H0, H0 + 4)},{ref('force')}+1),$B$9))")
    S("sel", '=IF($K$14="手动",0,IF(LEFT($K$14,2)="优化",5,0)+IFERROR(VALUE(RIGHT($K$14,1)),1))')
    S("rankMissing", f'=AND(COUNTIF($B${R_WINIT}:$M${R_WINIT},"—")+COUNTBLANK($B${R_WINIT}:$M${R_WINIT})>=12,'
                     f'COUNTA($B${R_WCUR}:$M${R_WCUR})=0)')
    S("manualBad", "=OR(" + ",".join(f"${MAN['from']}${H0 + j + 1}<${MAN['from']}${H0 + j}" for j in range(5)) + ")")
    assert srow[0] == H0 + len(names_order), "标量顺序不一致"

    # 向量：C0 / C4 / K / KT
    for j in range(NS):
        v, g = V9[j], rng(CLS["G"][j], H0, CLS_LAST)
        H(f"{v}{H0}", f"={N[0]}*IF({ref('startIdx')}=0,MAX(0,{SC[j]}$7),INDEX({g},{ref('startIdx')}))")
        H(f"{v}{H0 + 1}", f"=IF({ref('tgtIdx')}=0,0,{N[4]}*INDEX({g},{ref('tgtIdx')}))")
        H(f"{v}{H0 + 2}", f"=$B$14+{v}{H0}+{v}{H0 + 1}")
        H(f"{v}{H0 + 3}", (f"={v}{H0 + 2}+INDEX({rng(FL[0]['C'][j], H0, H0 + DATA_ROWS - 1)},${RK['i1']}${H0})"
                           f"+INDEX({rng(FL[1]['C'][j], H0, H0 + DATA_ROWS - 1)},${RK['i2']}${H0})"
                           f"+INDEX({rng(FL[2]['C'][j], H0, H0 + DATA_ROWS - 1)},${RK['i3']}${H0})"))
    # 等级表 / 当前熟练度 / 可练 / 推荐1累计可练
    for i, rk in enumerate(RANKS):
        H(f"{V13[i]}{H0}", rk)
    for s in range(NK):
        v = f'IF(${WC[s]}${R_WCUR}<>"",${WC[s]}${R_WCUR},${WC[s]}${R_WINIT})'
        H(f"{V13[s]}{H0 + 1}", f'=IF({v}="×",0,IFERROR(MATCH({v},{RANKROW},0),2))')
        H(f"{V13[s]}{H0 + 2}", f'=IF({v}="×",0,1)')
        tcol = rng(CLS["T"][s], H0, CLS_LAST)
        H(f"{V13[s]}{H0 + 3}", f"=INDEX({tcol},{Z(ref('startIdx'))})")
        for k in range(1, 5):
            H(f"{V13[s]}{H0 + 3 + k}", f"={V13[s]}{H0 + 2 + k}+INDEX({tcol},{Z(ref(f'P{k}'))})")

    # ------------------------------------------------------------------
    # 职业辅助表（行5~154 对应职业表第3~152行，行155为全0空行）
    # ------------------------------------------------------------------
    elig_lo = [LO[1], LO[2], LO[3], LO[4]]
    for i in range(DATA_ROWS):
        r, cr = H0 + i, 3 + i
        nm, mn, av = f"${CLS['name']}{r}", f"${CLS['min']}{r}", f"${CLS['avail']}{r}"
        H(f"{CLS['name']}{r}", f'=IF({CLASS_Q}!$C{cr}="","",{CLASS_Q}!$C{cr})')
        H(f"{CLS['min']}{r}", f"=N({CLASS_Q}!$D{cr})")
        H(f"{CLS['avail']}{r}", f'={CLASS_Q}!$E{cr}&""')
        H(f"{CLS['low']}{r}", f'=IF({av}="靠后",1,0)')
        for j in range(NS):
            H(f"{CLS['G'][j]}{r}", f'=IF({nm}="",0,MAX(0,{SC[j]}$7+N({CLASS_Q}!{DC[j]}{cr})))')
        for s in range(NK):
            key, L = SKILL_KEYS[s], len(SKILL_KEYS[s])
            H(f"{CLS['T'][s]}{r}", f'=IF({nm}="",0,IF(ISNUMBER(SEARCH("{key}",{CLASS_Q}!${C_USE}{cr})),1,0))')
            for tag, col in (("RM", C_MAIN), ("RS", C_SEL)):
                t = f"{CLASS_Q}!${col}{cr}"
                p = f'SEARCH("{key}",{t})'
                H(f"{CLS[tag][s]}{r}", (f'=IF({nm}="",0,IFERROR(MATCH(MID({t},{p}+{L},1)'
                                         f'&IF(MID({t},{p}+{L + 1},1)="+","+",""),{RANKROW},0),0))'))
            cur = f"{V13[s]}${H0 + 1}"
            H(f"{CLS['NM'][s]}{r}", f"=IF({CLS['RM'][s]}{r}>{cur},1,0)")
            H(f"{CLS['SU'][s]}{r}", f"=IF({CLS['RS'][s]}{r}>{cur},1,0)")
        rs_row = f"{CLS['RS'][0]}{r}:{CLS['RS'][-1]}{r}"
        H(f"{CLS['SA'][0] if isinstance(CLS['SA'], list) else CLS['SA']}{r}", f'=IF(COUNTIF({rs_row},">0")>0,1,0)')
        H(f"{CLS['SS']}{r}", f"=IF(SUMPRODUCT(({rs_row}>0)*({rs_row}<={CURROW}))>0,1,0)")
        ok_avail = f'OR({av}="是",{av}="靠后")'
        for k in range(3):
            el = CLS["EL"][k]
            H(f"{el}{r}", (f'=AND({nm}<>"",IF({ref("force")}={k + 1},{nm}=$C$10,'
                           f'OR({nm}=$C$9,AND({ok_avail},{mn}<=MAX($B$9,{elig_lo[k]})))))'))
            H(f"{CLS['CN'][k]}{r}", f'=IF({el}{r},COUNTIF({el}${H0}:{el}{r},TRUE),"")')
        for k in range(4):
            el = CLS["ELU"][k]
            H(f"{el}{r}", f'=AND({nm}<>"",OR({nm}=$C$9,AND({ok_avail},{mn}<=MAX($B$9,{elig_lo[k]}))))')
            H(f"{CLS['CNU'][k]}{r}", f'=IF({el}{r},COUNTIF({el}${H0}:{el}{r},TRUE),"")')

    # 各窗口候选列表
    n_eff = [N[1], N[2], f'({N[3]}+IF($C$10="",{N[4]},0))']
    for k in range(3):
        cn = rng(CLS["CN"][k], H0, CLS_LAST)
        for i in range(DATA_ROWS):
            r, p = H0 + i, i + 1
            ix = f"{FL[k]['idx']}{r}"
            H(ix, f"=IFERROR(MATCH({p},{cn},0),0)")
            H(f"{FL[k]['name']}{r}", f'=IF({ix}=0,"",INDEX({NAMES},{ix}))')
            for j in range(NS):
                H(f"{FL[k]['C'][j]}{r}", f"=IF({ix}=0,0,INDEX({rng(CLS['G'][j], H0, CLS_LAST)},{ix})*{n_eff[k]})")
    for k in range(4):
        cn = rng(CLS["CNU"][k], H0, CLS_LAST)
        for i in range(DATA_ROWS):
            H(f"{UL[k]}{H0 + i}", f"=IFERROR(MATCH({i + 1},{cn},0),0)")

    # ------------------------------------------------------------------
    # 组合枚举（每个区间一个职业）
    # ------------------------------------------------------------------
    header(ws, 4, list(CMB.keys()), start_col=ws[f"{CMB['i1']}1"].column)
    m1, m2, m3, M = ref("m1"), ref("m2"), ref("m3"), ref("M")
    si = ref("startIdx")
    for i in range(COMBO_ROWS):
        r = H0 + i
        ok = f"{i}<{M}"
        c = {k: f"{v}{r}" for k, v in CMB.items()}
        H(c["i1"], f"=IF({ok},INT({i}/({m2}*{m3}))+1,1)")
        H(c["i2"], f"=IF({ok},MOD(INT({i}/{m3}),{m2})+1,1)")
        H(c["i3"], f"=IF({ok},MOD({i},{m3})+1,1)")
        H(c["id1"], f"=IF({N[1]}>0,INDEX({FLI[0]},{c['i1']}),{si})")
        H(c["id2"], f"=IF({N[2]}>0,INDEX({FLI[1]},{c['i2']}),{c['id1']})")
        H(c["id3"], f"=IF({N[3]}>0,INDEX({FLI[2]},{c['i3']}),{c['id2']})")
        ids = [si, c["id1"], c["id2"], c["id3"]]
        terms = []
        u = trow(si)
        for k in range(1, 4):
            terms.append(f"IF(AND({N[k]}>0,{ids[k]}<>{ids[k - 1]},{ids[k]}<>{si}),{badc(ids[k], u)},0)")
            u = f"{u}+{trow(ids[k])}"
        tg = ref("tgtIdx")
        terms.append(f'IF(AND({N[4]}>0,$C$10<>"",{tg}<>{c["id3"]},{tg}<>{si}),{badc(tg, u)},0)')
        H(c["bad"], f"=IF({ok},{'+'.join(terms)},0)")
        id4 = f'IF($C$10<>"",{tg},{c["id3"]})'
        H(c["low"], (f"=IF({ok},({N[1]}>0)*INDEX({LOWC},{Z(c['id1'])})+({N[2]}>0)*INDEX({LOWC},{Z(c['id2'])})"
                     f"+({N[3]}>0)*INDEX({LOWC},{Z(c['id3'])})+({N[4]}>0)*INDEX({LOWC},{Z(id4)}),0)"))
        fin = (f"({START}+INT(({KROW}+INDEX({FLC[0]},{c['i1']},0)+INDEX({FLC[1]},{c['i2']},0)"
               f"+INDEX({FLC[2]},{c['i3']},0))/100))")
        mt, sh, tt = met_short_total(fin)
        H(c["met"], f"=IF({ok},{mt},0)")
        H(c["short"], f"=IF({ok},{sh},0)")
        H(c["total"], f"=IF({ok},{tt},0)")
        H(c["score"], score(f"AND({ok},{ref('listOK')})", c["bad"], c["low"], c["met"], c["short"], c["total"], i))

    # ------------------------------------------------------------------
    # 中途换职优化（基于推荐1，在某个区间内再换一次职业）
    # ------------------------------------------------------------------
    header(ws, 4, list(REF.keys()), start_col=ws[f"{REF['k']}1"].column)
    Rn = ref("R")
    for i in range(REF_ROWS):
        r = H0 + i
        ok = f"{i}<{Rn}"
        c = {k: f"{v}{r}" for k, v in REF.items()}
        k_ = c["k"]
        H(k_, f"=IF({ok},MATCH({i},{OROW},1),1)")
        H(c["j"], f"=IF({ok},INT(({i}-INDEX({OROW},{k_}))/INDEX({NNROW},{k_}))+1,1)")
        H(c["x"], f"=IF({ok},MOD({i}-INDEX({OROW},{k_}),INDEX({NNROW},{k_}))+1,1)")
        H(c["B"], f"=IF({ok},INDEX({ULMAT},{c['j']},{k_}),0)")
        P = f"INDEX({PROW},{k_})"
        hd = f"INDEX({HEADROW},{k_})"
        H(c["valid"], (f"=AND({ok},{ref('validT')},{c['B']}>0,{c['B']}<>{P},{c['x']}>=$B$13,"
                       f"INDEX({NW},{k_})-{c['x']}>=$B$13,"
                       f"OR({hd}=0,{ref('tgtMin')}<=INDEX({EW},{k_})+{c['x']}))"))
        u = f"INDEX({UCUM},{k_}+1-{hd},0)"
        H(c["bad"], f"=IF({c['valid']},{ref('badT')}+IF({c['B']}={si},0,{badc(c['B'], u)}),0)")
        H(c["low"], f"=IF({c['valid']},{ref('lowT')}+INDEX({LOWC},{Z(c['B'])}),0)")
        fin = (f"({START}+INT(({KTROW}+{c['x']}*(INDEX({GMAT},{Z(c['B'])},0)-INDEX({GMAT},{Z(P)},0)))/100))")
        mt, sh, tt = met_short_total(fin)
        H(c["met"], f"=IF({c['valid']},{mt},0)")
        H(c["short"], f"=IF({c['valid']},{sh},0)")
        H(c["total"], f"=IF({c['valid']},{tt},0)")
        H(c["score"], score(c["valid"], c["bad"], c["low"], c["met"], c["short"], c["total"], i))
        # 同一区间同一职业（连续的 x 行）只保留换职等级最优的一行，避免 Top5 被同职业不同等级占满
        sc_full = rng(REF["score"], H0, H0 + REF_ROWS - 1)
        b0 = f"({i}-{c['x']}+2)"
        b1 = f"MIN({REF_ROWS},{b0}+INDEX({NNROW},{k_})-1)"
        H(c["best"], (f"=IF({c['score']}>=MAX(INDEX({sc_full},{b0}):INDEX({sc_full},{b1})),"
                      f"{c['score']},-1E+12-{i})"))

    # ------------------------------------------------------------------
    # 排名：行5~9 推荐1~5，行10~14 优化1~5；每条路线展开为 6 段（起始等级 FROM / 职业序号 IDX）
    # ------------------------------------------------------------------
    header(ws, 4, list(k for k in RK.keys() if isinstance(RK[k], str)), start_col=ws[f"{RK['pos']}1"].column)
    csc = rng(CMB["score"], H0, H0 + COMBO_ROWS - 1)
    rsc = rng(REF["best"], H0, H0 + REF_ROWS - 1)
    base_r = H0
    for rr in range(10):
        r = H0 + rr
        is_opt = rr >= 5
        rank = rr % 5 + 1
        c = {k: f"${v}${r}" for k, v in RK.items() if isinstance(v, str)}
        src, sc_ = (REF, rsc) if is_opt else (CMB, csc)
        top = H0 + (REF_ROWS if is_opt else COMBO_ROWS) - 1
        H(c["pos"][1:].replace("$", ""), f"=MATCH(LARGE({sc_},{rank}),{sc_},0)")
        H(c["valid"].replace("$", ""), f"=INDEX({sc_},{c['pos']})>-1E+11")
        for key in ("met", "short", "total", "bad", "low"):
            H(c[key].replace("$", ""), f"=INDEX({rng(src[key], H0, top)},{c['pos']})")
        FROM = [f"${v}${r}" for v in RK["FROM"]]
        IDX = [f"${v}${r}" for v in RK["IDX"]]
        if not is_opt:
            for key in ("i1", "i2", "i3"):
                H(c[key].replace("$", ""), f"=INDEX({rng(CMB[key], H0, top)},{c['pos']})")
            H(IDX[0].replace("$", ""), f"={si}")
            for k in (1, 2, 3):
                H(IDX[k].replace("$", ""), f"=INDEX({rng(CMB[f'id{k}'], H0, top)},{c['pos']})")
            H(IDX[4].replace("$", ""), f'=IF($C$10<>"",{ref("tgtIdx")},{IDX[3]})')
            H(IDX[5].replace("$", ""), f"={IDX[4]}")
            for k in range(5):
                H(FROM[k].replace("$", ""), f"={E[k]}")
            H(FROM[5].replace("$", ""), "=999")
        else:
            for key in ("k", "x", "B"):
                H(c[key].replace("$", ""), f"=INDEX({rng(REF[key], H0, top)},{c['pos']})")
            k_ = c["k"]
            H(c["P"].replace("$", ""), f"=INDEX({PROW},{k_})")
            H(c["head"].replace("$", ""), f"=INDEX({HEADROW},{k_})")
            H(c["split"].replace("$", ""),
              f"=INDEX({EW},{k_})+IF({c['head']}=1,{c['x']},INDEX({NW},{k_})-{c['x']})")
            H(c["first"].replace("$", ""), f"=IF({c['head']}=1,{c['B']},{c['P']})")
            H(c["second"].replace("$", ""), f"=IF({c['head']}=1,{c['P']},{c['B']})")
            bIDX = [f"${v}${base_r}" for v in RK["IDX"]]
            bFROM = [f"${v}${base_r}" for v in RK["FROM"]]
            for i in range(6):
                prev = max(i - 1, 0)
                H(IDX[i].replace("$", ""),
                  f"=IF({i}<{k_},{bIDX[i]},IF({i}={k_},{c['first']},IF({i}={k_}+1,{c['second']},{bIDX[prev]})))")
                H(FROM[i].replace("$", ""),
                  f"=IF({i}<{k_},{bFROM[i]},IF({i}={k_},{bFROM[min(i, 4)]},IF({i}={k_}+1,{c['split']},{bFROM[prev]})))")
        NAME = [f"${v}${r}" for v in RK["NAME"]]
        NN_ = [f"${v}${r}" for v in RK["N"]]
        for i in range(6):
            H(NAME[i].replace("$", ""), f'=IF({IDX[i]}=0,{"$C$9" if i == 0 else chr(34) * 2},INDEX({NAMES},{IDX[i]}))')
            nxt = FROM[i + 1] if i < 5 else "999"
            H(NN_[i].replace("$", ""), f"=MAX(0,MIN($B$10,{nxt})-MAX($B$9,{FROM[i]}))")

    # ------------------------------------------------------------------
    # 可见区：③ 推荐 / ③+ 优化
    # ------------------------------------------------------------------
    def route_text(r):
        FROM = [f"${v}${r}" for v in RK["FROM"]]
        NAME = [f"${v}${r}" for v in RK["NAME"]]
        NN_ = [f"${v}${r}" for v in RK["N"]]
        parts = [f'IF({NN_[0]}>0,"Lv"&MAX($B$9,{FROM[0]})&" "&{NAME[0]}&" → ","")']
        for i in range(1, 6):
            parts.append(f'IF(AND({NN_[i]}>0,OR(MAX($B$9,{FROM[i]})=$B$9,{NAME[i]}<>{NAME[i - 1]})),'
                         f'"Lv"&MAX($B$9,{FROM[i]})&" "&{NAME[i]}&" → ","")')
        return "&".join(parts) + '&"Lv"&$B$10&" 完成"'

    for block, (rsec, rhdr, r0, title) in enumerate(
            ((R_BSEC, R_BHDR, R_B0, "③ 推荐路线 Top5（每个区间一个职业；排序：熟练度可行 ＞ 达成项数 ＞ 不含靠后职业 ＞ 缺口最小 ＞ 属性总和最大）"),
             (R_OSEC, R_OHDR, R_O0, "③+ 中途换职优化 Top5（在推荐1的某个区间中途再换一次职业，枚举换哪个职业、第几级换）"))):
        section(ws, rsec, title, LAST_VIS)
        header(ws, rhdr, ["方案", "达成项", "缺口合计"] + STATS + ["路线（LvX 职业 = 从该等级起在此职业升级）", "熟练度", "备注"])
        ws.merge_cells(f"O{rhdr}:{LAST_VIS}{rhdr}")
        for i in range(5):
            r, hr = r0 + i, H0 + block * 5 + i
            c = {k: f"${v}${hr}" for k, v in RK.items() if isinstance(v, str)}
            ws.cell(r, 1, ("优化" if block else "推荐") + str(i + 1))
            ws[f"B{r}"] = f'=IF({c["valid"]},{c["met"]}&"/"&{ref("tgtCnt")},"—")'
            ws[f"C{r}"] = f'=IF({c["valid"]},{c["short"]},"")'
            for j in range(NS):
                if block == 0:
                    ws[f"{SC[j]}{r}"] = (f'=IF({c["valid"]},{SC[j]}$9+INT(({V9[j]}${H0 + 2}'
                                         f'+INDEX({rng(FL[0]["C"][j], H0, H0 + DATA_ROWS - 1)},{c["i1"]})'
                                         f'+INDEX({rng(FL[1]["C"][j], H0, H0 + DATA_ROWS - 1)},{c["i2"]})'
                                         f'+INDEX({rng(FL[2]["C"][j], H0, H0 + DATA_ROWS - 1)},{c["i3"]}))/100),"")')
                else:
                    g = rng(CLS["G"][j], H0, CLS_LAST)
                    ws[f"{SC[j]}{r}"] = (f'=IF({c["valid"]},{SC[j]}$9+INT(({V9[j]}${H0 + 3}+{c["x"]}*'
                                         f'(INDEX({g},{Z(c["B"])})-INDEX({g},{Z(c["P"])})))/100),"")')
            none_txt = ('"无可行组合"' if i == 0 else '"—"')
            ws[f"M{r}"] = f"=IF({c['valid']},{route_text(hr)},{none_txt})"
            ws[f"N{r}"] = (f'=IF({c["valid"]},IF({c["bad"]}=0,"✓ 熟练度可行","✗ "&{c["bad"]}&"项熟练度无法练成")'
                           f'&IF({c["low"]}>0,"·含靠后职业",""),"")')
            ws.merge_cells(f"O{r}:{LAST_VIS}{r}")
            if block == 0:
                ws[f"O{r}"] = "基准路线（下方「优化」在此基础上中途换职）" if i == 0 else None
            else:
                b = {k: f"${RK[k]}${H0}" for k in ("met", "short", "total")}
                ws[f"O{r}"] = (f'=IF({c["valid"]},"对比推荐1：达成"&TEXT({c["met"]}-{b["met"]},"+0;-0;0")'
                               f'&"，缺口"&TEXT({c["short"]}-{b["short"]},"+0;-0;0")'
                               f'&"，属性总和"&TEXT({c["total"]}-{b["total"]},"+0;-0;0"),"")')
            style_range(ws, f"A{r}:{LAST_VIS}{r}", align=center)
            for col in "MNO":
                ws[f"{col}{r}"].alignment = left_nowrap
            ws[f"A{r}"].font = f_bold

    # ------------------------------------------------------------------
    # 当前显示路线 / 手动路线：辅助 6 段
    # ------------------------------------------------------------------
    sel = ref("sel")
    rk_rng = lambda col: rng(col, H0, H0 + 9)
    for j in range(6):
        r = H0 + j
        H(f"{SEG['from']}{r}", f"=IF({sel}=0,${MAN['from']}${r},INDEX({rk_rng(RK['FROM'][j])},{sel}))")
        H(f"{SEG['idx']}{r}", f"=IF({sel}=0,${MAN['idx']}${r},INDEX({rk_rng(RK['IDX'][j])},{sel}))")
        H(f"{SEG['name']}{r}", f"=IF({sel}=0,${MAN['name']}${r},INDEX({rk_rng(RK['NAME'][j])},{sel}))")
    for j in range(6):
        r, mr = H0 + j, R_M0 + j
        if j == 0:
            H(f"{MAN['from']}{r}", "=$B$9")
            H(f"{MAN['name']}{r}", "=$C$9")
            H(f"{MAN['idx']}{r}", f"={si}")
        else:
            H(f"{MAN['from']}{r}", f'=IF($B{mr}="",999,$B{mr})')
            H(f"{MAN['name']}{r}", f'=IF($B{mr}="","",$C{mr})')
            H(f"{MAN['idx']}{r}", f"=IFERROR(MATCH({MAN['name']}{r},{NAMES},0),0)")
    for G_ in (SEG, MAN):
        for j in range(6):
            r = H0 + j
            nxt = f"{G_['from']}{r + 1}" if j < 5 else "999"
            H(f"{G_['n']}{r}", f"=MAX(0,MIN($B$10,{nxt})-MAX($B$9,{G_['from']}{r}))")
            H(f"{G_['prev']}{r}", f"={si}" if j == 0 else f"={G_['idx']}{r - 1}")
            for s in range(NK):
                tcol = rng(CLS["T"][s], H0, CLS_LAST)
                if j == 0:
                    H(f"{G_['U'][s]}{r}", f"=INDEX({tcol},{Z(si)})")
                else:
                    H(f"{G_['U'][s]}{r}", f"={G_['U'][s]}{r - 1}+INDEX({tcol},{Z(G_['idx'] + str(r - 1))})")
            ix = Z(f"{G_['idx']}{r}")
            pieces, sel_pieces, scan_terms = [], [], []
            scan = f"{G_['scan']}{r}"
            for s in range(NK):
                can = f"({G_['U'][s]}{r}*{V13[s]}${H0 + 2}>0)"
                nm_ = rng(CLS["NM"][s], H0, CLS_LAST)
                su_ = rng(CLS["SU"][s], H0, CLS_LAST)
                rm_ = rng(CLS["RM"][s], H0, CLS_LAST)
                rs_ = rng(CLS["RS"][s], H0, CLS_LAST)
                pieces.append(f'IF(INDEX({nm_},{ix})=1,"{SKILLS[s]}"&INDEX({RANKROW},1,INDEX({rm_},{ix}))'
                              f'&IF({can},"(可练) ","(✗无法练) "),"")')
                # 选一：有能练的就只列能练的；一个都练不了才全部标 ✗
                sel_pieces.append(f'IF(AND(INDEX({su_},{ix})=1,OR({scan}=FALSE,{can})),'
                                  f'"{SKILLS[s]}"&INDEX({RANKROW},1,INDEX({rs_},{ix}))'
                                  f'&IF({can},"(可练) ","(✗) "),"")')
                scan_terms.append(f"AND(INDEX({su_},{ix})=1,{can})")
            H(scan, "=OR(" + ",".join(scan_terms) + ")")
            H(f"{G_['piece']}{r}", ("=" + "&".join(pieces)
                                    + f'&IF(AND(INDEX({SAC},{ix})=1,INDEX({SSC},{ix})=0),"｜选一："&'
                                    + "&".join(sel_pieces) + ',"")'))

    def req_cells(r, G_, hr, n_ref):
        """N=所需熟练度，O=熟练度检查，P=解锁条件"""
        ix = Z(f"${G_['idx']}${hr}")
        m = f"INDEX({class_col(C_MAIN)},{ix})"
        s_ = f"INDEX({class_col(C_SEL)},{ix})"
        m_none, s_none = f'OR({m}="",{m}="—")', f'OR({s_}="",{s_}="—")'
        idx = f"${G_['idx']}${hr}"
        ws[f"N{r}"] = (f'=IF({n_ref}=0,"",IF({idx}=0,"",IF(AND({m_none},{s_none}),"无",'
                       f'IF({m_none},"",{m})&IF({s_none},"",IF({m_none},"","；")&"选一："&{s_}))))')
        piece = f"${G_['piece']}${hr}"
        ws[f"O{r}"] = (f'=IF({n_ref}=0,"",IF({idx}=0,"⚠ 职业不在职业表",'
                       f'IF(OR({idx}=${G_["prev"]}${hr},{idx}={si}),"—（已在该职业）",'
                       f'IF({piece}="","✓ 满足",IF(ISNUMBER(SEARCH("✗",{piece})),"✗ 无法满足：","需练：")'
                       f'&IF(LEFT({piece},1)="｜",MID({piece},2,999),{piece})))))')
        cond = f"INDEX({class_col(C_COND)},{ix})"
        unl = f"INDEX({class_col(C_UNLOCK)},{ix})"
        ws[f"P{r}"] = (f'=IF({n_ref}=0,"",IF({idx}=0,"",TRIM(SUBSTITUTE({cond},CHAR(10)," "))'
                       f'&IF(OR({unl}="",{unl}="—"),"","；"&{unl})))')

    # ④ 当前显示路线
    section(ws, R_SSEC, "④ 当前显示路线：几级转什么职业、在哪个职业升几级、需要什么熟练度（由上方「显示路线」切换）", LAST_VIS)
    header(ws, R_SHDR, ["阶段", "起始等级", "职业", "结束等级", "升级次数"])
    ws.merge_cells(f"F{R_SHDR}:M{R_SHDR}")
    header(ws, R_SHDR, ["说明"], start_col=6)
    header(ws, R_SHDR, ["所需熟练度", "熟练度检查（可练=之前的职业能练）", "解锁条件"], start_col=14)
    for j in range(6):
        r, hr = R_S0 + j, H0 + j
        n = f"${SEG['n']}${hr}"
        ws.cell(r, 1, f"阶段{j + 1}")
        ws[f"B{r}"] = f'=IF({n}>0,MAX($B$9,${SEG["from"]}${hr}),"")'
        ws[f"C{r}"] = f'=IF({n}>0,${SEG["name"]}${hr},"")'
        ws[f"D{r}"] = f'=IF({n}>0,B{r}+{n},"")'
        ws[f"E{r}"] = f'=IF({n}>0,{n},"")'
        prev_name = "$C$9" if j == 0 else f"${SEG['name']}${hr - 1}"
        ws.merge_cells(f"F{r}:M{r}")
        ws[f"F{r}"] = (f'=IF({n}>0,"Lv"&B{r}&IF(B{r}=$B$9,IF(C{r}=$C$9," 以【"&C{r}&"】起步"," 立即转职为【"&C{r}&"】"),'
                       f'IF(C{r}={prev_name}," 继续【"&C{r}&"】"," 转职为【"&C{r}&"】"))'
                       f'&"，在该职业升 "&{n}&" 级 → Lv"&D{r},"")')
        req_cells(r, SEG, hr, n)
        style_range(ws, f"A{r}:{LAST_VIS}{r}", align=center)
        for col in "FNOP":
            ws[f"{col}{r}"].alignment = left
        ws[f"A{r}"].font = f_bold
        ws.row_dimensions[r].height = 30

    # 最终属性对比
    section(ws, R_FSEC, "最终属性对比（当前显示路线，目标等级时）", LAST_VIS)
    header(ws, R_FHDR, ["项目", "", ""] + STATS)
    for r, t in zip(range(R_F0, R_F0 + 4), ["路线结果", "目标", "差值", "判定"]):
        ws.cell(r, 1, t).font = f_bold
    for j in range(NS):
        s = SC[j]
        ws[f"{s}{R_F0}"] = f'=IFERROR(INDEX({s}${R_T0}:{s}${R_TLAST},$B$10-$B$9+1),"")'
        ws[f"{s}{R_F0 + 1}"] = f'=IF({s}$10="","",{s}$10)'
        ws[f"{s}{R_F0 + 2}"] = f'=IF(OR({s}{R_F0 + 1}="",{s}{R_F0}=""),"",{s}{R_F0}-{s}{R_F0 + 1})'
        ws[f"{s}{R_F0 + 3}"] = f'=IF({s}{R_F0 + 2}="","—",IF({s}{R_F0 + 2}>=0,"✓","✗ 差"&-{s}{R_F0 + 2}))'
    ws.merge_cells(f"B{R_F0 + 3}:C{R_F0 + 3}")
    ws[f"B{R_F0 + 3}"] = f'=COUNTIF(D{R_F0 + 3}:L{R_F0 + 3},"✓")&" / "&{ref("tgtCnt")}&" 项达成"'
    ws[f"B{R_F0 + 3}"].font = f_bold
    style_range(ws, f"A{R_F0}:L{R_F0 + 3}", align=center)

    # ⑤ 手动路线
    section(ws, R_MSEC, "⑤ 手动路线（「显示路线」选「手动」时生效：自己填几级转什么职业，下方逐级表会按此计算）", LAST_VIS)
    header(ws, R_MHDR, ["段", "转职等级", "职业"])
    ws.merge_cells(f"D{R_MHDR}:M{R_MHDR}")
    header(ws, R_MHDR, ["说明"], start_col=4)
    header(ws, R_MHDR, ["所需熟练度", "熟练度检查（可练=之前的职业能练）", "解锁条件"], start_col=14)
    for j in range(6):
        r, hr = R_M0 + j, H0 + j
        ws.cell(r, 1, f"段{j + 1}").font = f_bold
        ws.merge_cells(f"D{r}:M{r}")
        if j == 0:
            ws[f"B{r}"] = "=$B$9"
            ws[f"C{r}"] = "=$C$9"
            ws[f"D{r}"] = "起点（自动取上方起点等级/职业）"
            style_range(ws, f"B{r}:C{r}", fill=fill_grey, align=center)
        else:
            style_range(ws, f"B{r}:C{r}", font=f_input, fill=fill_input, align=center)
            ws[f"D{r}"] = "在此等级转职为该职业；不用的段留空" if j == 1 else None
        req_cells(r, MAN, hr, f"${MAN['n']}${hr}")
        style_range(ws, f"A{r}:A{r}", align=center)
        style_range(ws, f"D{r}:M{r}", font=f_note, align=left_nowrap)
        style_range(ws, f"N{r}:{LAST_VIS}{r}", align=left)
        ws.row_dimensions[r].height = 30
    ws[f"B{R_M0 + 1}"], ws[f"C{R_M0 + 1}"] = 20, "弓箭手"
    ws[f"B{R_M0 + 2}"], ws[f"C{R_M0 + 2}"] = 35, "狙击手"

    # ⑥ 逐级成长表
    section(ws, R_TSEC, "⑥ 逐级成长表（当前显示路线；绿色=该级属性+1，黄色行=转职）", LAST_VIS)
    header(ws, R_THDR, ["等级", "阶段", "职业"] + STATS + ["备注"])
    seg_from = rng(SEG["from"], H0, H0 + 5)
    for i in range(TL_ROWS):
        r = R_T0 + i
        ws[f"A{r}"] = f'=IF(OR($B$9="",$B$10=""),"",IF($B$9+{i}<=$B$10,$B$9+{i},""))'
        ws[f"B{r}"] = f'=IF($A{r}="","",MATCH($A{r},{seg_from},1))'
        ws[f"C{r}"] = f'=IF($A{r}="","",INDEX({rng(SEG["name"], H0, H0 + 5)},$B{r}))'
        H(f"{TLH['idx']}{r}", f'=IF($A{r}="",0,INDEX({rng(SEG["idx"], H0, H0 + 5)},$B{r}))')
        for j in range(NS):
            cum = TLH["cum"][j]
            g = rng(CLS["G"][j], H0, CLS_LAST)
            if i == 0:
                H(f"{cum}{r}", f'=IF($A{r}="","",0)')
            else:
                H(f"{cum}{r}", (f'=IF($A{r}="","",{cum}{r - 1}+IF(${TLH["idx"]}{r - 1}=0,MAX(0,{SC[j]}$7),'
                                f'INDEX({g},${TLH["idx"]}{r - 1})))'))
            ws[f"{SC[j]}{r}"] = f'=IF($A{r}="","",{SC[j]}$9+INT(($B$14+{cum}{r})/100))'
        if i == 0:
            ws[f"M{r}"] = (f'=IF($A{r}="","",IF($C{r}=$C$9,"起点：【"&$C{r}&"】",'
                           f'"起点：【"&$C$9&"】 ★ 立即转职 → 【"&$C{r}&"】"))')
        else:
            ws[f"M{r}"] = (f'=IF($A{r}="","",IF($C{r}<>$C{r - 1},"★ 转职 → 【"&$C{r}&"】","")'
                           f'&IF($A{r}=$B$10,"（目标等级）",""))')
        style_range(ws, f"A{r}:M{r}", align=center)
        ws[f"M{r}"].alignment = left_nowrap

    # ------------------------------------------------------------------
    # 警告
    # ------------------------------------------------------------------
    ws.merge_cells(f"A15:{LAST_VIS}15")
    tmin, elv = ref("tgtMin"), ref("entryLv")
    ws["A15"] = (
        '=IF(OR($B$9="",$B$10=""),"⚠ 请填写起点等级和目标等级；",IF($B$10<=$B$9,"⚠ 目标等级必须大于起点等级；",""))'
        f'&IF($B$10-$B$9>{TL_ROWS - 1},"⚠ 等级跨度超过{TL_ROWS - 1}，逐级表只显示前{TL_ROWS}级；","")'
        f'&IF({si}=0,"⚠ 起点职业不在职业表中（补正按0计）；","")'
        f'&IF(AND($C$10<>"",{ref("tgtIdx")}=0),"⚠ 目标职业不在职业表中；","")'
        f'&IF(AND({ref("tgtIdx")}>0,{elv}<{tmin}),"⚠ 目标职业最低转职等级为Lv"&{tmin}&"，但按目标等级只能在Lv"&{elv}&"转入；","")'
        f'&IF({M}>{COMBO_ROWS},"⚠ 候选组合"&{M}&"个，超过{COMBO_ROWS}，只搜索了前{COMBO_ROWS}个（可把用不到的职业设为「否」）；","")'
        f'&IF({Rn}>{REF_ROWS},"⚠ 中途换职候选"&{Rn}&"个，超过{REF_ROWS}，只搜索了前{REF_ROWS}个；","")'
        f'&IF(NOT({ref("listOK")}),"⚠ 某个阶段没有可选职业，请检查职业表「纳入推荐」；","")'
        f'&IF(NOT({ref("validT")}),"⚠ 没有可行的推荐路线；","")'
        f'&IF(AND({ref("validT")},{ref("badT")}>0),"⚠ 所有推荐路线都有练不成的熟练度，目标职业可能无法转入（请检查②熟练度或更换目标职业）；","")'
        f'&IF(AND({sel}=0,{ref("manualBad")}),"⚠ 手动路线的转职等级需从上到下递增、中间不要空行；","")'
        f'&IF({ref("rankMissing")},"⚠ 该角色没有熟练度数据，按E计算，建议在②填写当前熟练度；","")'
        f'&IF(COUNTIF(O{R_S0}:O{R_S0 + 5},"*✗*")>0,"⚠ 当前显示路线有无法练成的熟练度（见④）；","")'
    )
    ws["A15"].font = f_warn
    ws["A15"].alignment = left_nowrap

    # ------------------------------------------------------------------
    # 条件格式 / 数据验证
    # ------------------------------------------------------------------
    green = PatternFill("solid", fgColor="C6EFCE")
    red = PatternFill("solid", fgColor="FFC7CE")
    yellow = PatternFill("solid", fgColor="FFF2CC")
    gfont, rfont = Font(color="006100", bold=True), Font(color="9C0006", bold=True)
    cf = ws.conditional_formatting
    cf.add(f"D{R_T0 + 1}:L{R_TLAST}", FormulaRule(
        formula=[f"AND(ISNUMBER(D{R_T0 + 1}),ISNUMBER(D{R_T0}),D{R_T0 + 1}>D{R_T0})"], fill=green, font=gfont))
    cf.add(f"A{R_T0}:C{R_TLAST}", FormulaRule(formula=[f'LEFT($M{R_T0},1)="★"'], fill=yellow))
    cf.add(f"M{R_T0}:M{R_TLAST}", FormulaRule(formula=[f'ISNUMBER(SEARCH("★",$M{R_T0}))'], fill=yellow))
    r3 = R_F0 + 3
    cf.add(f"D{r3}:L{r3}", FormulaRule(formula=[f'LEFT(D{r3},1)="✓"'], fill=green, font=gfont))
    cf.add(f"D{r3}:L{r3}", FormulaRule(formula=[f'LEFT(D{r3},1)="✗"'], fill=red, font=rfont))
    cf.add(f"D{R_F0 + 2}:L{R_F0 + 2}", CellIsRule(operator="lessThan", formula=["0"], font=rfont))
    for r0 in (R_B0, R_O0):
        rows = f"{r0}:{r0 + 4}"
        cf.add(f"D{r0}:L{r0 + 4}", FormulaRule(
            formula=[f'AND(D$10<>"",ISNUMBER(D{r0}),D{r0}<D$10)'], fill=red, font=Font(color="9C0006")))
        cf.add(f"B{r0}:B{r0 + 4}", FormulaRule(formula=[f'B{r0}=({ref("tgtCnt")}&"/"&{ref("tgtCnt")})'], fill=green, font=gfont))
        cf.add(f"N{r0}:N{r0 + 4}", FormulaRule(formula=[f'LEFT(N{r0},1)="✓"'], fill=green, font=gfont))
        cf.add(f"N{r0}:N{r0 + 4}", FormulaRule(formula=[f'LEFT(N{r0},1)="✗"'], fill=red, font=rfont))
        del rows
    for r0 in (R_S0, R_M0):
        rr = f"O{r0}:O{r0 + 5}"
        cf.add(rr, FormulaRule(formula=[f'ISNUMBER(SEARCH("✗",O{r0}))'], fill=red, font=rfont))
        cf.add(rr, FormulaRule(formula=[f'LEFT(O{r0},2)="需练"'], fill=yellow))
        cf.add(rr, FormulaRule(formula=[f'LEFT(O{r0},1)="✓"'], fill=green, font=gfont))

    dv_char = DataValidation(type="list", formula1=CHAR_NAMES, allow_blank=True)
    dv_class = DataValidation(type="list", formula1=CLASS_NAMES, allow_blank=True)
    dv_route = DataValidation(type="list", formula1='"推荐1,推荐2,推荐3,推荐4,推荐5,优化1,优化2,优化3,优化4,优化5,手动"')
    dv_lv = DataValidation(type="whole", operator="between", formula1="1", formula2="200", allow_blank=True)
    dv_acc = DataValidation(type="whole", operator="between", formula1="0", formula2="99")
    dv_rank = DataValidation(type="list", formula1='"' + ",".join(RANKS + ["×"]) + '"', allow_blank=True)
    for dv in (dv_char, dv_class, dv_route, dv_lv, dv_acc, dv_rank):
        ws.add_data_validation(dv)
    dv_char.add("B4")
    dv_class.add("C9:C10")
    dv_class.add(f"C{R_M0 + 1}:C{R_M0 + 5}")
    dv_route.add("K14")
    dv_lv.add("B9:B10")
    dv_lv.add(f"B{R_M0 + 1}:B{R_M0 + 5}")
    dv_lv.add("E14:H14")
    dv_acc.add("B14")
    dv_lv.add("B13")
    dv_rank.add(f"B{R_WCUR}:M{R_WCUR}")

    ws.freeze_panes = "A7"
    ws.column_dimensions.group("Q", CL(LAST_HELPER), hidden=True)
    return ws


def main():
    wb = Workbook()  # 默认第一个 sheet 作为主页面
    n_char, char_cols = build_char_sheet(wb)
    n_class, class_cols = build_class_sheet(wb)
    build_main(wb, char_cols, class_cols)
    wb.calculation.fullCalcOnLoad = True
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print(f"saved {OUT}: {n_char} 角色, {n_class} 职业")


if __name__ == "__main__":
    main()
